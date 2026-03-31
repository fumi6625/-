"""
大阪レストラン検索 GUI ツール (Tkinter)

使い方: python search_gui.py [--file restaurants.xlsx]

Excelファイルからレストランデータを読み込み、
ジャンル・最寄駅・予算・席数でインタラクティブに検索します。
"""
import argparse
import os
import sys
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from tkinter.font import Font

try:
    import pandas as pd
    import openpyxl
except ImportError:
    print("必要なパッケージをインストール中... pip install pandas openpyxl")
    os.system(f"{sys.executable} -m pip install pandas openpyxl -q")
    import pandas as pd
    import openpyxl


# ──────────────────────────────────────────────────────────────────
# データ読み込み
# ──────────────────────────────────────────────────────────────────

SHEET_NAMES = ["食べログ", "ホットペッパーグルメ"]
COL_NAMES = [
    "店名", "最寄駅", "住所", "食べ物のジャンル",
    "コース料理", "飲み放題", "席数", "個室", "大体の予算（円）", "リンク",
]

BUDGET_OPTIONS = ["上限なし", "〜¥2,000", "〜¥3,000", "〜¥4,000", "〜¥5,000", "〜¥6,000"]
BUDGET_VALUES  = [0, 2000, 3000, 4000, 5000, 6000]


def load_data(filepath: str) -> dict[str, pd.DataFrame]:
    """Excelファイルからシートごとにデータを読み込む"""
    result = {}
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    for sheet_name in SHEET_NAMES:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        data = []
        headers = None
        for row in ws.iter_rows(values_only=True):
            if headers is None:
                headers = [str(c) if c else "" for c in row]
                continue
            if all(c is None for c in row):
                continue
            data.append(row)
        if data:
            df = pd.DataFrame(data, columns=headers)
            # 席数・予算を数値に
            for num_col in ["席数", "大体の予算（円）"]:
                if num_col in df.columns:
                    df[num_col] = pd.to_numeric(df[num_col], errors="coerce").fillna(0).astype(int)
            result[sheet_name] = df
    wb.close()
    return result


# ──────────────────────────────────────────────────────────────────
# GUI アプリ
# ──────────────────────────────────────────────────────────────────

class RestaurantSearchApp:

    def __init__(self, root: tk.Tk, data: dict[str, pd.DataFrame]):
        self.root = root
        self.data = data
        self.root.title("大阪レストラン検索")
        self.root.geometry("1400x750")
        self.root.configure(bg="#F0F4F8")

        self._build_ui()

    def _build_ui(self):
        # ─── タイトルバー ───────────────────────────────
        title_frame = tk.Frame(self.root, bg="#1F497D", height=50)
        title_frame.pack(fill="x")
        title_frame.pack_propagate(False)
        tk.Label(
            title_frame, text="🍽  大阪レストラン検索",
            font=("Helvetica", 16, "bold"),
            bg="#1F497D", fg="white",
        ).pack(side="left", padx=15, pady=10)

        # ─── 検索条件 ───────────────────────────────────
        search_frame = tk.LabelFrame(
            self.root, text="  検索条件  ", font=("Helvetica", 10, "bold"),
            bg="#F0F4F8", fg="#1F497D", relief="groove", bd=2, padx=10, pady=8,
        )
        search_frame.pack(fill="x", padx=15, pady=(10, 5))

        # 条件行1
        row1 = tk.Frame(search_frame, bg="#F0F4F8")
        row1.pack(fill="x", pady=4)

        # データソース
        tk.Label(row1, text="データソース:", bg="#F0F4F8", width=12, anchor="e").pack(side="left")
        self.var_source = tk.StringVar(value="両方")
        source_options = ["両方"] + list(self.data.keys())
        ttk.Combobox(
            row1, textvariable=self.var_source,
            values=source_options, state="readonly", width=20,
        ).pack(side="left", padx=(4, 20))

        # ジャンル
        tk.Label(row1, text="ジャンル:", bg="#F0F4F8", width=9, anchor="e").pack(side="left")
        all_genres = self._collect_unique_values("食べ物のジャンル")
        self.var_genre = tk.StringVar(value="すべて")
        genre_cb = ttk.Combobox(
            row1, textvariable=self.var_genre,
            values=["すべて"] + all_genres, width=18,
        )
        genre_cb.pack(side="left", padx=(4, 20))

        # 最寄駅
        tk.Label(row1, text="最寄駅:", bg="#F0F4F8", width=8, anchor="e").pack(side="left")
        all_stations = self._collect_unique_values("最寄駅")
        self.var_station = tk.StringVar(value="すべて")
        station_cb = ttk.Combobox(
            row1, textvariable=self.var_station,
            values=["すべて"] + all_stations, state="readonly", width=16,
        )
        station_cb.pack(side="left", padx=(4, 20))

        # 条件行2
        row2 = tk.Frame(search_frame, bg="#F0F4F8")
        row2.pack(fill="x", pady=4)

        # 予算上限
        tk.Label(row2, text="予算上限:", bg="#F0F4F8", width=12, anchor="e").pack(side="left")
        self.var_budget = tk.StringVar(value="上限なし")
        ttk.Combobox(
            row2, textvariable=self.var_budget,
            values=BUDGET_OPTIONS, state="readonly", width=12,
        ).pack(side="left", padx=(4, 20))

        # 最低席数（スライダー付き）
        tk.Label(row2, text="最低席数:", bg="#F0F4F8", width=9, anchor="e").pack(side="left")
        self.var_seats = tk.IntVar(value=0)
        self.lbl_seats_val = tk.Label(row2, text="指定なし", bg="#F0F4F8", width=10, anchor="w")
        scale = ttk.Scale(
            row2, from_=0, to=100, orient="horizontal",
            variable=self.var_seats, length=150,
            command=lambda v: self.lbl_seats_val.config(
                text="指定なし" if int(float(v)) == 0 else f"{int(float(v))}席以上"
            ),
        )
        scale.pack(side="left", padx=(4, 5))
        self.lbl_seats_val.pack(side="left", padx=(0, 20))

        # コース・飲み放題・個室フィルタ
        self.var_course = tk.BooleanVar(value=False)
        self.var_drink = tk.BooleanVar(value=False)
        self.var_private = tk.BooleanVar(value=False)
        tk.Checkbutton(row2, text="コースあり", variable=self.var_course, bg="#F0F4F8").pack(side="left", padx=5)
        tk.Checkbutton(row2, text="飲み放題あり", variable=self.var_drink, bg="#F0F4F8").pack(side="left", padx=5)
        tk.Checkbutton(row2, text="個室あり", variable=self.var_private, bg="#F0F4F8").pack(side="left", padx=5)

        # 検索ボタン
        btn_frame = tk.Frame(search_frame, bg="#F0F4F8")
        btn_frame.pack(fill="x", pady=6)
        tk.Button(
            btn_frame, text="  🔍  検索  ", command=self._search,
            bg="#2E75B6", fg="white", font=("Helvetica", 11, "bold"),
            relief="flat", cursor="hand2", padx=10, pady=6,
        ).pack(side="left", padx=5)
        tk.Button(
            btn_frame, text="  ↺  リセット  ", command=self._reset,
            bg="#7F7F7F", fg="white", font=("Helvetica", 10),
            relief="flat", cursor="hand2", padx=8, pady=6,
        ).pack(side="left", padx=5)
        self.lbl_result_count = tk.Label(
            btn_frame, text="", bg="#F0F4F8", fg="#1F497D",
            font=("Helvetica", 10, "bold"),
        )
        self.lbl_result_count.pack(side="left", padx=20)

        # ─── 結果テーブル ────────────────────────────────
        table_frame = tk.Frame(self.root, bg="#F0F4F8")
        table_frame.pack(fill="both", expand=True, padx=15, pady=(0, 10))

        cols = ("出典", "店名", "最寄駅", "ジャンル", "コース", "飲み放題", "席数", "個室", "予算（円）", "住所")
        self.tree = ttk.Treeview(table_frame, columns=cols, show="headings", height=20)

        col_widths = [90, 200, 90, 120, 65, 75, 55, 65, 90, 350]
        for col, w in zip(cols, col_widths):
            self.tree.heading(col, text=col, command=lambda c=col: self._sort_column(c))
            self.tree.column(col, width=w, minwidth=40)

        # スクロールバー
        vsb = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)

        # 行色タグ
        self.tree.tag_configure("tabelog", background="#FFF5EE")
        self.tree.tag_configure("hotpepper", background="#FFF0F0")
        self.tree.tag_configure("alt_tabelog", background="#FFE8D6")
        self.tree.tag_configure("alt_hotpepper", background="#FFE0E0")

        # ダブルクリックで URL を開く
        self.tree.bind("<Double-1>", self._on_double_click)

        # 初期表示（全件）
        self._search()

    # ──────────────────────────────────────────
    # 検索ロジック
    # ──────────────────────────────────────────

    def _search(self):
        source = self.var_source.get()
        genre  = self.var_genre.get()
        station = self.var_station.get()
        budget_label = self.var_budget.get()
        budget_max = BUDGET_VALUES[BUDGET_OPTIONS.index(budget_label)]
        seats_min = int(self.var_seats.get())
        need_course  = self.var_course.get()
        need_drink   = self.var_drink.get()
        need_private = self.var_private.get()

        # 対象シートを決定
        if source == "両方":
            targets = list(self.data.items())
        else:
            targets = [(source, self.data[source])] if source in self.data else []

        results = []
        for sheet_name, df in targets:
            filtered = df.copy()

            if genre != "すべて":
                filtered = filtered[filtered["食べ物のジャンル"].fillna("").str.contains(genre, na=False)]
            if station != "すべて":
                filtered = filtered[filtered["最寄駅"].fillna("").str.contains(station, na=False)]
            if budget_max > 0:
                filtered = filtered[
                    (filtered["大体の予算（円）"] == 0) | (filtered["大体の予算（円）"] <= budget_max)
                ]
            if seats_min > 0:
                filtered = filtered[
                    (filtered["席数"] == 0) | (filtered["席数"] >= seats_min)
                ]
            if need_course:
                filtered = filtered[filtered.get("コース料理", pd.Series()).fillna("").str.contains("有")]
            if need_drink:
                filtered = filtered[filtered.get("飲み放題", pd.Series()).fillna("").str.contains("有")]
            if need_private:
                filtered = filtered[filtered.get("個室", pd.Series()).fillna("").str.contains("有")]

            for _, row in filtered.iterrows():
                results.append((sheet_name, row))

        # テーブル更新
        self.tree.delete(*self.tree.get_children())
        self._current_data = []  # URL保持用

        for i, (sheet_name, row) in enumerate(results):
            short_name = "食べログ" if "食べ" in sheet_name else "ホットペッパー"
            budget_val = int(row.get("大体の予算（円）", 0) or 0)
            seats_val  = int(row.get("席数", 0) or 0)
            values = (
                short_name,
                row.get("店名", ""),
                row.get("最寄駅", ""),
                row.get("食べ物のジャンル", ""),
                row.get("コース料理", ""),
                row.get("飲み放題", ""),
                f"{seats_val}" if seats_val > 0 else "",
                row.get("個室", ""),
                f"¥{budget_val:,}" if budget_val > 0 else "",
                row.get("住所", ""),
            )
            tag = ("tabelog" if "食べ" in sheet_name else "hotpepper") if i % 2 == 0 \
                  else ("alt_tabelog" if "食べ" in sheet_name else "alt_hotpepper")
            self.tree.insert("", "end", values=values, tags=(tag,))
            self._current_data.append(row.get("リンク", ""))

        self.lbl_result_count.config(text=f"検索結果: {len(results)} 件")

    def _reset(self):
        self.var_source.set("両方")
        self.var_genre.set("すべて")
        self.var_station.set("すべて")
        self.var_budget.set("上限なし")
        self.var_seats.set(0)
        self.lbl_seats_val.config(text="指定なし")
        self.var_course.set(False)
        self.var_drink.set(False)
        self.var_private.set(False)
        self._search()

    def _on_double_click(self, event):
        item = self.tree.selection()
        if not item:
            return
        idx = self.tree.index(item[0])
        url = self._current_data[idx] if idx < len(self._current_data) else ""
        if url and url.startswith("http"):
            import webbrowser
            webbrowser.open(url)
        else:
            messagebox.showinfo("リンク", "URLが設定されていません。")

    def _sort_column(self, col: str):
        items = [(self.tree.set(k, col), k) for k in self.tree.get_children("")]
        items.sort()
        for idx, (_, k) in enumerate(items):
            self.tree.move(k, "", idx)

    def _collect_unique_values(self, col_name: str) -> list[str]:
        vals = set()
        for df in self.data.values():
            if col_name in df.columns:
                vals.update(df[col_name].dropna().astype(str).unique())
        return sorted(v for v in vals if v)


# ──────────────────────────────────────────────────────────────────
# エントリーポイント
# ──────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="大阪レストラン検索 GUI")
    parser.add_argument("--file", default="restaurants.xlsx",
                        help="Excelファイルパス（デフォルト: restaurants.xlsx）")
    args = parser.parse_args()

    filepath = args.file
    if not os.path.exists(filepath):
        # ファイル選択ダイアログ
        root = tk.Tk()
        root.withdraw()
        filepath = filedialog.askopenfilename(
            title="レストランデータExcelを選択",
            filetypes=[("Excel ファイル", "*.xlsx *.xlsm"), ("すべて", "*.*")],
        )
        root.destroy()
        if not filepath:
            print("ファイルが選択されませんでした。")
            return

    print(f"読み込み中: {filepath}")
    data = load_data(filepath)

    if not data:
        print("データが見つかりませんでした。")
        return

    total = sum(len(df) for df in data.values())
    print(f"読み込み完了: {total} 件")

    root = tk.Tk()
    app = RestaurantSearchApp(root, data)
    root.mainloop()


if __name__ == "__main__":
    main()
